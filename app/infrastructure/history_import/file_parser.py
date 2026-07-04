import csv
from pathlib import Path

from openpyxl import load_workbook


def _build_column_map(headers: list[str], columns: dict) -> dict:
    mapping = {}
    for col_key, col_def in columns.items():
        aliases = [a.lower() for a in col_def["header_aliases"]]
        for idx, header in enumerate(headers):
            if str(header).strip().lower() in aliases:
                mapping[col_key] = idx
                break
    return mapping


def _is_empty_row(row) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)


def _coerce_value(value, col_type: str):
    """Normalize values so CSV (all strings) and xlsx (native types) return consistent types."""
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if col_type == "float":
        try:
            return float(s.replace(",", ""))
        except ValueError:
            return s
    return s


def _map_row(raw_row: list, column_map: dict, columns: dict, source_row: int) -> dict:
    result = {"_source_row": source_row}
    for col_key, col_def in columns.items():
        idx = column_map.get(col_key)
        raw = raw_row[idx] if idx is not None and idx < len(raw_row) else None
        col_type = col_def.get("type", "string")
        value = _coerce_value(raw, col_type)
        if value is None and "default" in col_def:
            value = col_def["default"]
        result[col_key] = value
    return result


def parse_xlsx(file_path: str, profile: dict) -> list[dict]:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    skip_rows = profile.get("skip_rows", 0)
    rows_iter = ws.iter_rows(values_only=True)

    headers = None
    column_map = {}
    results = []
    for i, row in enumerate(rows_iter, start=1):
        if i <= skip_rows:
            continue
        if headers is None:
            headers = [str(h).strip() if h is not None else "" for h in row]
            column_map = _build_column_map(headers, profile["columns"])
            continue
        if _is_empty_row(row):
            continue
        results.append(_map_row(row, column_map, profile["columns"], i))

    wb.close()
    return results


def parse_csv(file_path: str, profile: dict) -> list[dict]:
    skip_rows = profile.get("skip_rows", 0)
    with open(file_path, encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        headers = None
        column_map = {}
        results = []
        for i, row in enumerate(reader, start=1):
            if i <= skip_rows:
                continue
            if headers is None:
                headers = [h.strip() for h in row]
                column_map = _build_column_map(headers, profile["columns"])
                continue
            if _is_empty_row(row):
                continue
            results.append(_map_row(row, column_map, profile["columns"], i))
    return results


def parse_file(file_path: str, profile: dict) -> list[dict]:
    ext = Path(file_path).suffix.lower()
    if ext == ".xlsx":
        return parse_xlsx(file_path, profile)
    elif ext == ".csv":
        return parse_csv(file_path, profile)
    else:
        raise ValueError(f"Unsupported file format: {ext}")
