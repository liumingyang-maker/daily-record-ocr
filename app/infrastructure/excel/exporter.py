"""Excel exporter — generates .xlsx from job data."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    for col_idx, _ in enumerate(headers, 1):
        max_len = len(str(headers[col_idx - 1]))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)


def export_job_to_excel(job_id: int, data: dict) -> Path:
    """Create an xlsx file with 4 sheets and return its path."""
    from app.settings import EXPORTS_DIR

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Sheet 1: 生产记录汇总
    ws1 = wb.active
    ws1.title = "生产记录汇总"
    _write_sheet(
        ws1,
        ["记录号", "时间", "客户", "产品", "颜色", "牌号"],
        [
            [
                r["记录号"],
                r["时间"],
                r["客户"],
                r["产品"],
                r["颜色"],
                r["牌号"],
            ]
            for r in data.get("records", [])
        ],
    )

    # Sheet 2: 配方明细
    ws2 = wb.create_sheet("配方明细")
    _write_sheet(
        ws2,
        ["记录号", "序号", "物料名称", "用量", "单位", "置信度"],
        [
            [
                m["记录号"],
                m["序号"],
                m["物料名称"],
                m["用量"],
                m["单位"],
                m["置信度"],
            ]
            for m in data.get("materials", [])
        ],
    )

    # Sheet 3: 识别审查
    ws3 = wb.create_sheet("识别审查")
    _write_sheet(
        ws3,
        ["记录号", "字段", "OCR结果", "MiMo结果", "最终值", "置信度", "来源", "需确认"],
        [
            [
                f["记录号"],
                f["字段"],
                f["OCR结果"],
                f["MiMo结果"],
                f["最终值"],
                f["置信度"],
                f["来源"],
                f["需确认"],
            ]
            for f in data.get("field_results", [])
        ],
    )

    # Sheet 4: 导入修正摘要
    ws4 = wb.create_sheet("导入修正摘要")
    _write_sheet(
        ws4,
        ["类型", "记录号", "字段", "旧值", "新值", "操作时间"],
        [
            [
                c["类型"],
                c["记录号"],
                c["字段"],
                c["旧值"],
                c["新值"],
                c["操作时间"],
            ]
            for c in data.get("corrections", [])
        ],
    )

    job_no = data.get("job_no", f"job_{job_id}")
    out_path = EXPORTS_DIR / f"{job_no}.xlsx"
    wb.save(str(out_path))
    return out_path
