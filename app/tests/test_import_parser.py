import os
import tempfile

import pytest
from openpyxl import Workbook

from app.infrastructure.history_import import load_import_profiles
from app.infrastructure.history_import.file_parser import parse_file


@pytest.fixture
def profile():
    profiles = load_import_profiles()
    return profiles["type_a_flat"]


@pytest.fixture
def xlsx_file(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "客户", "产品", "批号", "物料", "用量"])
    ws.append(["2025-01-01", "客户A", "产品X", "B001", "物料1", 10.5])
    ws.append(["2025-01-02", "客户B", "产品Y", "B002", "物料2", 20.0])
    path = tmp_path / "test.xlsx"
    wb.save(path)
    return str(path)


@pytest.fixture
def csv_file(tmp_path):
    path = tmp_path / "test.csv"
    with open(path, "w", encoding="utf-8") as f:
        f.write("日期,客户,产品,批号,物料,用量\n")
        f.write("2025-01-01,客户A,产品X,B001,物料1,10.5\n")
        f.write("2025-01-02,客户B,产品Y,B002,物料2,20.0\n")
    return str(path)


def test_parse_xlsx(xlsx_file, profile):
    rows = parse_file(xlsx_file, profile)
    assert len(rows) == 2
    assert rows[0]["date"] == "2025-01-01"
    assert rows[0]["customer"] == "客户A"
    assert rows[0]["usage"] == 10.5
    assert rows[0]["_source_row"] == 2
    assert rows[1]["_source_row"] == 3


def test_parse_csv(csv_file, profile):
    rows = parse_file(csv_file, profile)
    assert len(rows) == 2
    assert rows[0]["date"] == "2025-01-01"
    assert rows[0]["customer"] == "客户A"
    assert rows[0]["usage"] == "10.5"
    assert rows[0]["_source_row"] == 2


def test_skip_empty_rows(tmp_path, profile):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "客户", "产品", "批号", "物料", "用量"])
    ws.append(["2025-01-01", "客户A", "产品X", "B001", "物料1", 10.5])
    ws.append([None, None, None, None, None, None])
    ws.append(["2025-01-02", "客户B", "产品Y", "B002", "物料2", 20.0])
    ws.append([None, None, None, None, None, None])
    path = tmp_path / "empty_rows.xlsx"
    wb.save(path)
    rows = parse_file(str(path), profile)
    assert len(rows) == 2
    assert rows[0]["customer"] == "客户A"
    assert rows[1]["customer"] == "客户B"


def test_unsupported_format(tmp_path, profile):
    path = tmp_path / "test.txt"
    path.write_text("some text")
    with pytest.raises(ValueError, match="Unsupported file format"):
        parse_file(str(path), profile)


def test_unit_default(tmp_path, profile):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "客户", "产品", "批号", "物料", "用量"])
    ws.append(["2025-01-01", "客户A", "产品X", "B001", "物料1", 10.5])
    path = tmp_path / "default_unit.xlsx"
    wb.save(path)
    rows = parse_file(str(path), profile)
    assert rows[0]["unit"] == "kg"


def test_header_alias_matching(tmp_path, profile):
    wb = Workbook()
    ws = wb.active
    ws.append(["生产日期", "客户名称", "型号", "批次", "原料", "重量"])
    ws.append(["2025-01-01", "客户A", "产品X", "B001", "物料1", 10.5])
    path = tmp_path / "alias.xlsx"
    wb.save(path)
    rows = parse_file(str(path), profile)
    assert rows[0]["date"] == "2025-01-01"
    assert rows[0]["customer"] == "客户A"
    assert rows[0]["product"] == "产品X"
    assert rows[0]["batch_no"] == "B001"
    assert rows[0]["material_name"] == "物料1"
    assert rows[0]["usage"] == 10.5


def test_load_import_profiles():
    profiles = load_import_profiles()
    assert "type_a_flat" in profiles
    assert "type_b_template" in profiles
    p = profiles["type_a_flat"]
    assert "date" in p["columns"]
    assert p["columns"]["unit"]["default"] == "kg"
