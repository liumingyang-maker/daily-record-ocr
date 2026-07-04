"""Tests for Excel exporter."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.infrastructure.database.base import Base
from app.infrastructure.database.models import *  # noqa: F401,F403
from app.infrastructure.database.models import (
    Customer,
    FieldRecognitionResult,
    ManualCorrectionLog,
    ProductionRecord,
    RecordMaterialItem,
    RecognitionJob,
)


@pytest.fixture()
def db_engine():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    yield engine
    engine.dispose()


@pytest.fixture()
def db_session(db_engine):
    with Session(db_engine) as session:
        yield session


@pytest.fixture()
def patch_exports_dir(monkeypatch, tmp_path):
    exports_dir = tmp_path / "exports"
    exports_dir.mkdir()
    monkeypatch.setattr("app.settings.EXPORTS_DIR", exports_dir)
    return exports_dir


def _build_job_data(db_session):
    """Seed a minimal job with records, materials, field results, corrections."""
    from app.infrastructure.excel.exporter import export_job_to_excel

    customer = Customer(customer_code="C001", customer_name="客户A")
    db_session.add(customer)
    db_session.flush()

    job = RecognitionJob(job_no="JOB_TEST_001", status="CONFIRMED")
    db_session.add(job)
    db_session.flush()

    rec = ProductionRecord(
        job_id=job.id,
        record_index=1,
        time_raw="08:00",
        time_value="08:00",
        customer_id=customer.id,
        color_raw="红色",
        date_batch_no_raw="B001",
    )
    db_session.add(rec)
    db_session.flush()

    mat = RecordMaterialItem(
        record_id=rec.id,
        seq=1,
        material_name_standard="物料A",
        usage_value=10.5,
        unit_standard="kg",
        confidence=0.92,
    )
    db_session.add(mat)

    fr = FieldRecognitionResult(
        job_id=job.id,
        record_id=rec.id,
        field_key="material_name",
        field_label="物料名称",
        ocr_raw_text="物料A",
        mimo_raw_text="物料A",
        final_value="物料A",
        final_confidence=0.92,
        source="ocr",
        need_review=False,
    )
    db_session.add(fr)

    log = ManualCorrectionLog(
        record_id=rec.id,
        field_key="material_name",
        correction_type="OCR_ERROR",
        old_value="物料X",
        new_value="物料A",
    )
    db_session.add(log)
    db_session.commit()

    data = {
        "job_no": job.job_no,
        "records": [
            {
                "记录号": rec.record_index,
                "时间": rec.time_value,
                "客户": "客户A",
                "产品": "",
                "颜色": rec.color_raw,
                "牌号": rec.date_batch_no_raw,
            }
        ],
        "materials": [
            {
                "记录号": rec.record_index,
                "序号": mat.seq,
                "物料名称": mat.material_name_standard,
                "用量": mat.usage_value,
                "单位": mat.unit_standard,
                "置信度": mat.confidence,
            }
        ],
        "field_results": [
            {
                "记录号": rec.record_index,
                "字段": fr.field_label,
                "OCR结果": fr.ocr_raw_text,
                "MiMo结果": fr.mimo_raw_text,
                "最终值": fr.final_value,
                "置信度": fr.final_confidence,
                "来源": fr.source,
                "需确认": "是" if fr.need_review else "否",
            }
        ],
        "corrections": [
            {
                "类型": log.correction_type,
                "记录号": rec.record_index,
                "字段": log.field_key,
                "旧值": log.old_value,
                "新值": log.new_value,
                "操作时间": str(log.created_at),
            }
        ],
    }

    out_path = export_job_to_excel(job.id, data)
    return out_path


def test_export_creates_file_with_four_sheets(db_session, patch_exports_dir):
    out_path = _build_job_data(db_session)
    assert Path(out_path).exists()

    wb = load_workbook(str(out_path))
    assert wb.sheetnames == [
        "生产记录汇总",
        "配方明细",
        "识别审查",
        "导入修正摘要",
    ]

    ws1 = wb["生产记录汇总"]
    assert ws1.cell(1, 1).value == "记录号"
    assert ws1.cell(1, 1).font.bold is True
    assert ws1.cell(2, 1).value == 1
    assert ws1.cell(2, 3).value == "客户A"

    ws2 = wb["配方明细"]
    assert ws2.cell(2, 3).value == "物料A"
    assert ws2.cell(2, 4).value == 10.5

    ws3 = wb["识别审查"]
    assert ws3.cell(2, 3).value == "物料A"
    assert ws3.cell(2, 8).value == "否"

    ws4 = wb["导入修正摘要"]
    assert ws4.cell(2, 1).value == "OCR_ERROR"
    assert ws4.cell(2, 4).value == "物料X"
    assert ws4.cell(2, 5).value == "物料A"


def test_export_empty_job(db_session, patch_exports_dir):
    """Export a job with no records — sheets should have headers only."""
    from app.infrastructure.excel.exporter import export_job_to_excel

    job = RecognitionJob(job_no="JOB_EMPTY", status="CONFIRMED")
    db_session.add(job)
    db_session.commit()

    data = {"job_no": job.job_no}
    out_path = export_job_to_excel(job.id, data)
    assert Path(out_path).exists()

    wb = load_workbook(str(out_path))
    assert len(wb.sheetnames) == 4
    for ws in wb.worksheets:
        assert ws.max_row == 1  # header row only
