"""Config-driven Excel export, including existing fixed templates."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from lite_app.config import load_settings
from lite_app.storage import job_dir, load_job, load_result, save_job

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
_HEADER_FONT = Font(bold=True)


def _get_path(value: Any, path: str) -> Any:
    if not path:
        return value
    current = value
    for part in path.split("."):
        if isinstance(current, dict):
            current = current.get(part)
        else:
            return None
    return current


def _resolve_expression(
    expression: Any,
    *,
    root: dict[str, Any],
    item: Any,
    parent: Any,
    index: int,
    parent_index: int,
) -> Any:
    if not isinstance(expression, str):
        return expression
    if expression == "$index":
        return index
    if expression == "$parent_index":
        return parent_index
    if expression.startswith("$root."):
        return _get_path(root, expression[6:])
    if expression.startswith("$parent."):
        return _get_path(parent, expression[8:])
    if expression.startswith("literal:"):
        return expression[8:]
    value = _get_path(item, expression)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _iter_rows(root: dict[str, Any], table: dict[str, Any]):
    source = _get_path(root, str(table.get("source", "")))
    if not isinstance(source, list):
        return
    expand = str(table.get("expand", "")).strip()
    if not expand:
        for index, item in enumerate(source, start=1):
            yield item, None, index, index
        return

    for parent_index, parent in enumerate(source, start=1):
        children = _get_path(parent, expand)
        if not isinstance(children, list):
            continue
        for index, item in enumerate(children, start=1):
            yield item, parent, index, parent_index


def _prepare_workbook(config: dict[str, Any]):
    template_path = str(config.get("template_path", "")).strip()
    if template_path:
        path = load_settings().resolve_path(template_path)
        if path.exists():
            return load_workbook(path, keep_vba=bool(config.get("keep_vba", False)))
        raise FileNotFoundError(f"Configured Excel template does not exist: {path}")
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    return workbook


def _write_headers(worksheet, columns: list[dict[str, Any]], row: int) -> None:
    for column in columns:
        letter = str(column["column"]).upper()
        cell = worksheet[f"{letter}{row}"]
        cell.value = column.get("header", "")
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _auto_width(worksheet, columns: list[dict[str, Any]], start_row: int, end_row: int) -> None:
    for column in columns:
        letter = str(column["column"]).upper()
        column_index = column_index_from_string(letter)
        max_length = 0
        for row in worksheet.iter_rows(
            min_row=start_row,
            max_row=max(end_row, start_row),
            min_col=column_index,
            max_col=column_index,
        ):
            value = row[0].value
            if value is not None:
                max_length = max(max_length, len(str(value)))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 2, 8), 48
        )


def export_job(job_id: str) -> Path:
    settings = load_settings()
    config = settings.export.get("excel", {})
    if not isinstance(config, dict):
        raise ValueError("export.yaml must contain an 'excel' mapping")

    root = load_result(job_id)
    if root is None:
        raise ValueError("Recognition result does not exist")

    workbook = _prepare_workbook(config)

    cells = config.get("cells", [])
    if isinstance(cells, list):
        for mapping in cells:
            if not isinstance(mapping, dict):
                continue
            sheet_name = str(mapping["sheet"])
            worksheet = (
                workbook[sheet_name]
                if sheet_name in workbook.sheetnames
                else workbook.create_sheet(sheet_name)
            )
            value = _resolve_expression(
                mapping.get("value", ""),
                root=root,
                item=root,
                parent=None,
                index=1,
                parent_index=1,
            )
            worksheet[str(mapping["cell"])] = value

    tables = config.get("tables", [])
    if not isinstance(tables, list):
        raise ValueError("excel.tables must be a list")

    for table in tables:
        if not isinstance(table, dict):
            continue
        sheet_name = str(table.get("sheet", "识别结果"))
        worksheet = (
            workbook[sheet_name]
            if sheet_name in workbook.sheetnames
            else workbook.create_sheet(sheet_name)
        )
        columns = table.get("columns", [])
        if not isinstance(columns, list) or not columns:
            continue

        start_row = int(table.get("start_row", 1))
        include_header = bool(table.get("include_header", True))
        if include_header:
            _write_headers(worksheet, columns, start_row)
        current_row = start_row + 1 if include_header else start_row

        for item, parent, index, parent_index in _iter_rows(root, table) or []:
            for column in columns:
                letter = str(column["column"]).upper()
                value = _resolve_expression(
                    column.get("value", ""),
                    root=root,
                    item=item,
                    parent=parent,
                    index=index,
                    parent_index=parent_index,
                )
                worksheet[f"{letter}{current_row}"] = value
            current_row += 1

        if bool(table.get("auto_width", True)):
            _auto_width(worksheet, columns, start_row, current_row - 1)

    if not workbook.sheetnames:
        workbook.create_sheet("识别结果")

    output_name = str(config.get("output_name", "recognized-{job_id}.xlsx")).format(
        job_id=job_id
    )
    output_path = job_dir(job_id) / output_name
    workbook.save(output_path)

    job = load_job(job_id)
    job["export_file"] = output_path.name
    job["status"] = "EXPORTED"
    save_job(job)
    return output_path
